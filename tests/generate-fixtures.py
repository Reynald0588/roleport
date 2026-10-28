from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from docx import Document
from pypdf import PdfReader, PdfWriter
import shutil

ROOT = Path(r'C:\Users\oneJ\Documents\ChatGPT\个人网站\tests\fixtures')
ROOT.mkdir(parents=True, exist_ok=True)
pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))

# UTF-8 text and Markdown
(ROOT / 'sample-utf8.txt').write_text('TXT_MARKER_2026\n中文文本解析测试\nPython SQL 数据分析', encoding='utf-8')
(ROOT / 'sample-markdown.md').write_text('# MD_MARKER_2026\n\n- 中文简历测试\n- Solidity TypeScript\n', encoding='utf-8')
(ROOT / 'empty.txt').write_bytes(b'')

# DOCX
doc = Document()
doc.add_heading('DOCX_MARKER_2026', 0)
doc.add_paragraph('文档解析测试：负责社区运营与用户增长。')
doc.add_paragraph('技能：Solidity、TypeScript、Kubernetes。')
doc.save(ROOT / 'sample.docx')

# Text PDF
pdf_path = ROOT / 'sample-text.pdf'
c = canvas.Canvas(str(pdf_path), pagesize=A4)
c.setFont('STSong-Light', 12)
for i, line in enumerate([
    'PDF_MARKER_2026',
    '中文PDF测试：舆情监测与社区运营',
    '技能：Python SQL Kubernetes',
    'English reading and writing proficiency'
]):
    c.drawString(60, 780 - i * 24, line)
c.save()

# Two-column PDF
pdf_path = ROOT / 'sample-two-column.pdf'
c = canvas.Canvas(str(pdf_path), pagesize=A4)
c.setFont('STSong-Light', 11)
left = ['LEFT_COLUMN_MARKER', '教育经历', '渤海大学 食品科学与工程']
right = ['RIGHT_COLUMN_MARKER', '项目经历', 'Web3 空投项目实践']
for i, line in enumerate(left):
    c.drawString(55, 760 - i * 28, line)
for i, line in enumerate(right):
    c.drawString(320, 760 - i * 28, line)
c.save()

# Uppercase extension
shutil.copyfile(ROOT / 'sample-text.pdf', ROOT / 'sample-uppercase.PDF')

# Scanned/image-only PDF
img = Image.new('RGB', (1100, 720), 'white')
draw = ImageDraw.Draw(img)
font_cn = ImageFont.truetype(r'C:\Windows\Fonts\msyh.ttc', 36)
font_en = ImageFont.truetype(r'C:\Windows\Fonts\segoeui.ttf', 34)
draw.text((70, 70), 'SCANNED_PDF_MARKER_2026', fill='black', font=font_en)
draw.text((70, 150), '这是一份图片型扫描 PDF', fill='black', font=font_cn)
draw.text((70, 220), '页面没有可复制的文本层', fill='black', font=font_cn)
img_path = ROOT / '_scanned.png'
img.save(img_path)
c = canvas.Canvas(str(ROOT / 'sample-scanned.pdf'), pagesize=A4)
c.drawImage(ImageReader(str(img_path)), 0, 0, width=A4[0], height=A4[1])
c.save()
img_path.unlink()

# Encrypted PDF
base = ROOT / '_encrypted-base.pdf'
c = canvas.Canvas(str(base), pagesize=A4)
c.setFont('STSong-Light', 12)
c.drawString(60, 760, 'ENCRYPTED_PDF_MARKER_2026')
c.save()
reader = PdfReader(str(base))
writer = PdfWriter()
for page in reader.pages:
    writer.add_page(page)
writer.encrypt('secret')
with open(ROOT / 'sample-encrypted.pdf', 'wb') as stream:
    writer.write(stream)
base.unlink()

# Invalid / unsupported files
(ROOT / 'corrupt.pdf').write_text('This is not a PDF file.', encoding='utf-8')
(ROOT / 'fake.docx').write_text('This is not a DOCX ZIP archive.', encoding='utf-8')
(ROOT / 'sample.doc').write_text('Legacy Word document placeholder.', encoding='utf-8')
(ROOT / 'sample.xlsx').write_text('Unsupported spreadsheet placeholder.', encoding='utf-8')

print('FIXTURES_CREATED')
for path in sorted(ROOT.iterdir()):
    print(path.name, path.stat().st_size)
